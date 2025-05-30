import sqlite3
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from fpdf import FPDF
import os

def generate_binomes_from_db(niveau_nom, taille_groupe):
    # Récupération des étudiants depuis la base
    conn = sqlite3.connect("./db/scannerSYS.db")
    cursor = conn.cursor()
    cursor.execute("""
        SELECT e.code_apoge, e.nom, e.prenom
        FROM etudiant e
        JOIN niveau n ON e.id_niveau = n.id_niveau
        WHERE n.nom_niveau = ?
        ORDER BY e.nom, e.prenom
    """, (niveau_nom,))
    etudiants = cursor.fetchall()
    conn.close()

    if not etudiants:
        raise ValueError("Aucun étudiant trouvé pour ce niveau.")

    # Construction du DataFrame avec groupes
    table = []
    groupe_col = []
    groupe_num = 1
    for i, (code, nom, prenom) in enumerate(etudiants):
        table.append([code, nom, prenom])
        groupe_col.append(f"G{groupe_num}" if i % taille_groupe == 0 else "")
        if (i + 1) % taille_groupe == 0:
            groupe_num += 1

    df = pd.DataFrame(table, columns=["Code Apogee", "Nom", "Prenom"])
    df.insert(0, "Groupe", groupe_col)

    os.makedirs("outputs", exist_ok=True)
    excel_path = f"outputs/binomes_{niveau_nom}.xlsx"
    pdf_path = f"outputs/binomes_{niveau_nom}.pdf"
    temp_excel = "outputs/temp.xlsx"

    # === Export Excel ===
    df.to_excel(temp_excel, index=False)
    wb = load_workbook(temp_excel)
    ws = wb.active
    border_style = Border(bottom=Side(border_style="thick", color="000000"))
    for i in range(2, ws.max_row + 1):
        cur = ws.cell(row=i, column=1).value
        nxt = ws.cell(row=i + 1, column=1).value if i + 1 <= ws.max_row else None
        if cur != nxt:
            for cell in ws[i]:
                cell.border = border_style
    wb.save(excel_path)
    os.remove(temp_excel)

    # === Export PDF ===
    class PDF(FPDF):
        def header(self):
            self.set_font("Arial", "B", 14)
            self.cell(0, 10, f"Binômes {niveau_nom}", ln=True, align="C")
            self.ln(5)

    pdf = PDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    col_widths = [30, 60, 60]
    headers = ["Groupe", "Nom", "Prénom"]

    pdf.set_fill_color(0, 100, 0)
    pdf.set_text_color(255)
    pdf.set_font("Arial", "B", 11)
    for header, w in zip(headers, col_widths):
        pdf.cell(w, 10, header, border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Arial", "", 10)
    pdf.set_text_color(0)
    for i in range(0, len(df), taille_groupe):
        g = f"G{i // taille_groupe + 1}"
        for j in range(taille_groupe):
            if i + j < len(df):
                row = df.iloc[i + j]
                pdf.cell(col_widths[0], 10, g if j == 0 else "", border=1, align="C")
                pdf.cell(col_widths[1], 10, str(row["Nom"]), border=1)
                pdf.cell(col_widths[2], 10, str(row["Prenom"]), border=1)
                pdf.ln()

    pdf.output(pdf_path)

    return excel_path, pdf_path
