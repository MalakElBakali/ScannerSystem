import cv2
import numpy as np
import pytesseract
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
import os


def extract_table_with_ocr(image_path, output_excel_path):
    # Vérification de l'image
    img = cv2.imread(image_path)
    if img is None:
        raise FileNotFoundError(f"Image introuvable : {image_path}")

    # Prétraitement
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    _, binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)

    # Détection des lignes verticales et horizontales
    vertical_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, img.shape[0] // 30))
    vertical_lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, vertical_kernel, iterations=2)

    horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (img.shape[1] // 30, 1))
    horizontal_lines = cv2.morphologyEx(binary, cv2.MORPH_OPEN, horizontal_kernel, iterations=2)

    table_mask = cv2.add(vertical_lines, horizontal_lines)

    # Détection des cellules par contours
    contours, _ = cv2.findContours(table_mask, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    cells = [cv2.boundingRect(c) for c in contours if cv2.contourArea(c) > 500]
    cells = sorted(cells, key=lambda b: (b[1], b[0]))  # Trier par ligne, puis colonne

    # Regrouper les cellules par lignes
    rows = []
    current_row = []
    last_y = -100
    tolerance = 25

    for (x, y, w, h) in cells:
        if abs(y - last_y) > tolerance:
            if current_row:
                rows.append(current_row)
            current_row = [(x, y, w, h)]
            last_y = y
        else:
            current_row.append((x, y, w, h))
    if current_row:
        rows.append(current_row)

    # Création du fichier Excel
    wb = openpyxl.Workbook()
    ws = wb.active

    for i, row in enumerate(rows):
        row = sorted(row, key=lambda b: b[0])  # Tri horizontal
        for j, (x, y, w, h) in enumerate(row):
            # Extraire image cellule
            cell_img = img[y:y+h, x:x+w]
            cell_gray = cv2.cvtColor(cell_img, cv2.COLOR_BGR2GRAY)
            cell_thresh = cv2.threshold(cell_gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]

            # OCR pour la cellule
            config = "--psm 7"
            text = pytesseract.image_to_string(cell_thresh, config=config).strip()

            # Écriture dans Excel
            cell = ws.cell(row=i+1, column=j+1, value=text)
            col_letter = get_column_letter(j+1)
            ws.row_dimensions[i+1].height = h // 2
            ws.column_dimensions[col_letter].width = max(ws.column_dimensions[col_letter].width, w // 7)
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # Sauvegarde Excel
    wb.save(output_excel_path)
    print(f"✅ Tableau avec texte exporté dans : {output_excel_path}")


if __name__ == "__main__":
    # === Modifier ces chemins si nécessaire
    image_path = "./images/liste11.jpg"
    output_excel_path = "./outputs/output_with_text.xlsx"

    # Créer le dossier de sortie s’il n'existe pas
    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)

    # Exécuter la fonction
    extract_table_with_ocr(image_path, output_excel_path)
