import sqlite3
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
import os
from fpdf import FPDF


# === PARAMÈTRES ===
db_path = "./db/scannerSYS.db"
niveau_nom = "GINF2"        # Nom du niveau pour filtrer
taille_groupe = 2           # Binôme = 2
output_path = "./outputs/liste_binomes_styled.xlsx"

# === 1. Connexion à la base et récupération des étudiants ===
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

cursor.execute("""
SELECT e.code_apoge, e.nom, e.prenom
FROM etudiant e
JOIN niveau n ON e.id_niveau = n.id_niveau
WHERE n.nom_niveau = ?
ORDER BY e.nom, e.prenom
""", (niveau_nom,))

etudiants = cursor.fetchall()
conn.close()

# === 2. Création du tableau avec groupe
table = []
groupe_col = []
groupe_num = 1

for i, (code_apoge, nom, prenom) in enumerate(etudiants):
    table.append([code_apoge, nom, prenom])
    if i % taille_groupe == 1:
        groupe_col.append(f"G{groupe_num}")
        groupe_num += 1
    else:
        groupe_col.append("")

# === 3. Créer le DataFrame
df = pd.DataFrame(table, columns=["Code Apogée", "Nom", "Prénom"])
df.insert(0, "Groupe", groupe_col)

# === 4. Sauvegarde temporaire avant stylisation
os.makedirs("outputs", exist_ok=True)
temp_file = "outputs/temp_binomes.xlsx"
df.to_excel(temp_file, index=False)

# === 5. Ajout des bordures entre chaque groupe
wb = load_workbook(temp_file)
ws = wb.active
border_style = Border(bottom=Side(border_style="thick", color="000000"))

for i in range(taille_groupe + 1, ws.max_row + 1, taille_groupe):
    for cell in ws[i]:
        cell.border = border_style

# === 6. Sauvegarde finale
wb.save(output_path)
os.remove(temp_file)


class PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 14)
        self.cell(0, 10, "Liste des binômes", ln=True, align="C")
        self.ln(5)

def generate_pdf(df, pdf_path, niveau_nom):
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()

    # === TITRE DYNAMIQUE ===
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, f"Binômes {niveau_nom}", ln=True, align="C")
    pdf.ln(5)

    # === EN-TÊTE TABLEAU ===
    pdf.set_fill_color(0, 100, 0)
    pdf.set_text_color(255)
    pdf.set_font("Arial", "B", 11)
    col_widths = [30, 60, 60]
    headers = ["Groupe", "Nom", "Prénom"]
    for header, w in zip(headers, col_widths):
        pdf.cell(w, 10, header, border=1, align="C", fill=True)
    pdf.ln()

    # === CONTENU ===
    pdf.set_font("Arial", "", 10)
    pdf.set_text_color(0)
    for i in range(0, len(df), 2):
        groupe = f"G{i // 2 + 1}"
        row1 = df.iloc[i]
        row2 = df.iloc[i + 1] if i + 1 < len(df) else pd.Series(["", ""])

        pdf.cell(col_widths[0], 10, groupe, border=1, align="C")
        pdf.cell(col_widths[1], 10, str(row1["Nom"]), border=1)
        pdf.cell(col_widths[2], 10, str(row1["Prénom"]), border=1)
        pdf.ln()

        pdf.cell(col_widths[0], 10, "", border=1)
        pdf.cell(col_widths[1], 10, str(row2["Nom"]), border=1)
        pdf.cell(col_widths[2], 10, str(row2["Prénom"]), border=1)
        pdf.ln()

    pdf.output(pdf_path)
    print(f"📄 PDF binômes 2 lignes généré : {pdf_path}")

pdf_path = "./outputs/liste_binomes_styled.pdf"
generate_pdf(df, pdf_path, niveau_nom)
print(f"✅ Binômes générés depuis la base et exportés avec style : {output_path}")
